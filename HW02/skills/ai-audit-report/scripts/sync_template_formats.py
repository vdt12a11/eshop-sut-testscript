"""Preserve audit data cells while applying the retained AI-02 template."""
from pathlib import Path
import html, re, hashlib, json
from copy import deepcopy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

def rows(text):
    return [line for line in text.splitlines() if line.startswith('| AI-')]

def sync(root, template):
    root, template = Path(root), Path(template)
    path=root/'AI_Audit_Report.md'
    old=path.read_text(encoding='utf-8')
    raw=rows(old)
    assert raw
    d=Document(template)
    paras=[p.text for p in d.paragraphs]
    section4=re.search(r'^## 4\..*?\n(.*?)^## 5\.',old,re.M|re.S).group(1)
    counts={v:int(re.search(r'\| '+v+r' \| (\d+)',section4).group(1)) for v in ['VALID','INVALID','INCOMPLETE']}
    total=len(raw); reviewed=sum(counts.values())
    info=['NGUYỄN SƠN TÙNG','23120399','Chưa xác nhận','HW02','Xem ngày giờ nguyên văn tại từng tương tác','Codex','[x] Có  [ ] Không']
    for r,val in zip(d.tables[0].rows[1:],info): r.cells[1].text=val
    table=d.tables[1]
    pattern=deepcopy(table.rows[3]._tr)
    for r in list(table.rows)[1:]: table._tbl.remove(r._tr)
    header=table.rows[0]._tr.get_or_add_trPr()
    repeat=OxmlElement('w:tblHeader'); header.append(repeat)
    expected=[]
    for line in raw:
        values=[html.unescape(x.strip().replace('<br>','\n')) for x in line[1:-1].split('|')]
        assert len(values)==5
        expected.append(values)
        tr=deepcopy(pattern); table._tbl.append(tr)
        r=table.rows[-1]
        for c,value in zip(r.cells,values):
            # Retain template cell and paragraph properties; replace sample text.
            p=c.paragraphs[0]
            for x in list(c._tc):
                if x.tag==qn('w:p') and x is not p._p:c._tc.remove(x)
            p.text=value
            p.paragraph_format.keep_with_next=False
            p.paragraph_format.keep_together=False
        for el in list(tr.findall('./w:trPr/w:trHeight',tr.nsmap)):el.getparent().remove(el)
        for el in list(tr.findall('./w:trPr/w:cantSplit',tr.nsmap)):el.getparent().remove(el)
    summary=d.tables[2]
    summary.rows[1].cells[1].text=str(total)
    summary.rows[1].cells[2].text='Chưa đánh giá: '+str(total-reviewed)
    for r,v in zip(list(summary.rows)[2:],counts):
        r.cells[1].text=str(counts[v]);r.cells[2].text=(f'{counts[v]/reviewed:.2%}' if reviewed else 'Chưa tính')
    # Preserve the template's instructions and disclosure verbatim, without signing.
    note=d.paragraphs[19].insert_paragraph_before('Chưa có kết luận 80–150 chữ được sinh viên xác nhận.')
    for p in list(d.paragraphs):
        if p.text.startswith('________________'):p._element.getparent().remove(p._element)
    p=next(p for p in d.paragraphs if p.text.startswith('"[Test case'))
    p.insert_paragraph_before('Mẫu khai báo dưới đây chưa được sinh viên điền/xác nhận; không phải lời cam đoan đã ký.')
    sig=d.tables[3]
    sig.rows[0].cells[1].text=info[0];sig.rows[1].cells[1].text=info[1]
    for i in [2,4,5]:sig.rows[i].cells[1].text='Chưa xác nhận'
    sig.rows[6].cells[1].text='Chưa ký'
    decl='Tôi sử dụng công cụ AI cho các nhiệm vụ sau:\n- Đọc đề, đặc tả và lý thuyết; thiết kế EC, Domain Testing và BVA.\n- Tạo/chỉnh báo cáo, skill, script kiểm thử, tổng kết kết quả, bug report và nhật ký AI.'
    d.paragraphs[6].insert_paragraph_before(decl)
    scope='Phạm vi nhật ký: giữ nguyên prompt/output và thời gian đã ghi; chưa bao gồm toàn bộ bytes của mọi artifact lịch sử hoặc toàn bộ tool log. Mỗi hàng hiện ghi một tương tác; các file sinh ra được lưu snapshot để truy vết. Chưa có đánh giá riêng cho từng file artifact.'
    next(p for p in d.paragraphs if p.text.startswith('3. Bảng Audit')).insert_paragraph_before(scope)
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for run in p.runs:
                        run.font.name='Arial';run.font.size=Pt(10.5)
    out=root/'AI_Audit_Report.docx';d.save(out)
    reread=Document(out)
    assert [[c.text for c in r.cells] for r in list(reread.tables[1].rows)[1:]]==expected
    # Markdown follows the same sections/tables as Word, preserving encoded data rows.
    md=['# AI Audit Report — Mẫu 5 mục cho mỗi Artifact\n\n']
    for p in paras[:3]:md.append(p+'\n\n')
    md.extend([paras[4]+'\n\n',paras[5]+'\n\n',decl+'\n\n','## '+paras[6]+'\n\n'])
    def mdtable(t):
        rr=[[c.text.replace('\n','<br>').replace('|','&#124;') for c in r.cells] for r in t.rows]
        return '| '+' | '.join(rr[0])+' |\n| '+' | '.join(['---']*len(rr[0]))+' |\n'+''.join('| '+' | '.join(r)+' |\n' for r in rr[1:])+'\n'
    md.append(mdtable(d.tables[0]));md.append('## '+paras[7]+'\n\n')
    for p in paras[8:15]:md.append('- '+p+'\n')
    md.extend(['\n## '+paras[15]+'\n\n',scope+'\n\n','| (1) Prompt + Công cụ | (2) Output AI | (3) Verdict | (4) Lý do (ISTQB) | (5) Bản SV sửa |\n| --- | --- | --- | --- | --- |\n','\n'.join(raw)+'\n\n','## '+paras[16]+'\n\n',paras[17]+'\n\n',mdtable(summary),'## '+paras[18]+'\n\n',paras[19]+'\n\nChưa có kết luận được sinh viên xác nhận.\n\n','## '+paras[26]+'\n\n','Mẫu khai báo chưa được sinh viên điền/xác nhận:\n\n'+paras[27]+'\n\n','## Chữ ký\n\n',mdtable(sig),'## Tham khảo\n\n'])
    md.extend('- '+p+'\n' for p in paras[30:])
    md.append('\nNhật ký nguyên văn bổ sung: Appendix_A_Prompt_Log.md.\n')
    new=''.join(md);assert rows(new)==raw
    path.write_text(new,encoding='utf-8')
    (root/'generated_report_state.json').write_text(json.dumps({'sha256':hashlib.sha256(new.encode()).hexdigest(),'owner':'inline-audit-template'}),encoding='utf-8')
    print('Template synchronized; lossless dates/prompts/outputs:',total,'rows')

if __name__=='__main__':
    import argparse
    p=argparse.ArgumentParser();p.add_argument('--root',required=True);p.add_argument('--template',required=True)
    a=p.parse_args();sync(a.root,a.template)
